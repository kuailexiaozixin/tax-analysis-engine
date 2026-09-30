#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""《减免税政策代码目录》查询式子技能。

母技能主线里，这一层补的是**中央层级现行有效 / 已失效税收优惠的结构化目录**：
按税种、政策大类、减免性质代码穷举"某主体名下有哪些优惠"，这是五个在线检索源
给不了的——它们按关键词召回单份文件，给不出一个封闭、去重、带官方减免性质的
优惠全集。

数据来自国家税务总局官网每月定期发布的《减免税政策代码目录》xlsx（依据
国家税务总局公告 2015 年第 73 号，每月更新）。本技能不硬编码版本节点，直链现抓：

    抓「纳税服务」栏目页 → 定位「减免税政策代码目录」的 xlsx 直链
    → 与 data/sync/tax-preference/state.json 比对 → 变了才下载重建

同步（含版本 diff、下载守卫、内容 SHA1 复用）复用母技能 scripts/tax_sync.py；
定位失败/下载异常都不覆盖已有索引。全程 plain HTTP，不调用任何模型、不消耗额度。

用法：
    python preference.py sync            # 检查并按需下载重建索引
    python preference.py sync --check    # 只检查有无新版本，不下载
    python preference.py query 研发费用   # 离线查询（多词 AND）
    python preference.py query --code 01010503
    python preference.py query 小微 --type 增值税 --status 有效
    python preference.py list-types      # 列收入种类及条数
"""

import hashlib
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS_DIR = (Path(HERE).parent.parent / "scripts").resolve()
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

INDEX_PATH = Path(HERE) / "preference_index.json"

COLUMN_PAGE = "https://www.chinatax.gov.cn/chinatax/n810346/index.html"
HOME_PAGE = "https://www.chinatax.gov.cn/"
# 按顺序尝试的定位入口：栏目页是目录常驻落点，改版或临时不可达时回退首页。
ENTRY_PAGES = (COLUMN_PAGE, HOME_PAGE)
ANCHOR = "减免税政策代码目录"

# ── 文号 / 标题抽取（与官方目录"政策名称"列的两段式写法对齐）──────────────
RE_DOC_NO = re.compile(
    r"([\u4e00-\u9fa5A-Za-z]{2,12}?"
    r"(?:〔\s*\d{4}\s*〕|\[\s*\d{4}\s*\]|\(\s*\d{4}\s*\))"
    r"\s*第?\s*\d+\s*号)")
RE_ANNOUNCE = re.compile(r"([\u4e00-\u9fa5]{2,15}公告\s*\d{4}\s*年\s*第\s*\d+\s*号)")
RE_ORDER = re.compile(r"([\u4e00-\u9fa5]{2,12}令\s*第?[\u4e00-\u9fa5\d]+\s*号)")
RE_CUSTOM = re.compile(r"([\u4e00-\u9fa5]{2,12}\s*\d{4}\s*年\s*第\s*\d+\s*号)")
_PATTERNS = (RE_DOC_NO, RE_ANNOUNCE, RE_ORDER, RE_CUSTOM)


def _norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (datetime,)):
        return v.strftime("%Y/%m/%d")
    return re.sub(r"\s+", " ", str(v)).strip()


def extract_title(policy: str) -> str:
    m = re.search(r"《([^》]*)》", policy or "")
    return m.group(1).strip() if m else (policy or "").strip()


def extract_doc_no(policy: str) -> str:
    if not policy:
        return ""
    s = re.sub(r"\s+", " ", policy).strip()
    rest = re.sub(r"《[^》]*》", " ", s).strip()
    if rest and any(rx.search(rest) for rx in _PATTERNS):
        return re.sub(r"\s+", " ", rest)
    for rx in _PATTERNS:
        m = rx.search(s)
        if m:
            return re.sub(r"\s+", " ", m.group(1))
    if rest and len(rest) <= 40 and "号" in rest:
        return rest
    return ""


# ── xlsx → 索引 ────────────────────────────────────────────────────────────
def _fill_merged(ws, grid, start_row):
    nrows = len(grid)
    for rng in list(ws.merged_cells.ranges):
        r1, c1, r2, c2 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        if r2 < start_row:
            continue
        r1 = max(r1, start_row)
        if r1 - 1 >= nrows:
            continue
        val = grid[r1 - 1][c1 - 1] if c1 - 1 < len(grid[r1 - 1]) else None
        if val is None:
            continue
        for r in range(r1, min(r2, nrows) + 1):
            for c in range(c1, c2 + 1):
                if c - 1 < len(grid[r - 1]):
                    grid[r - 1][c - 1] = val


def _find_header(grid):
    for i, row in enumerate(grid):
        if sum(1 for c in row if _norm(c)) >= 5:
            return i
    raise RuntimeError("未找到表头行")


def _parse_sheet(ws, status):
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = _find_header(grid)
    _fill_merged(ws, grid, hi + 1)
    header = [_norm(c) for c in grid[hi]]
    ncols = max((j + 1 for j, h in enumerate(header) if h), default=0)
    header = header[:ncols]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    idx = {
        "收入种类": col("收入种类"),
        "大类": col("减免政策大类", "减免性质大类"),
        "小类": col("减免政策小类", "减免性质小类"),
        "代码": col("减免性质代码"),
        "政策名称": col("政策名称"),
        "有效期起": col("有效期起"),
        "有效期止": col("有效期止"),
        "优惠条款": col("优惠条款"),
        "减免项目名称": col("减免项目名称"),
        "关联政策": col("关联政策条款", "关联政策"),
    }
    recs = []
    for row in grid[hi + 1:]:
        vals = [_norm(c) for c in row[:ncols]]
        if not any(vals):
            continue
        name = vals[idx["政策名称"]] if idx["政策名称"] is not None else ""
        rec = {k: (vals[i] if i is not None and i < len(vals) else "")
               for k, i in idx.items()}
        rec["文件标题"] = extract_title(name)
        rec["文号"] = extract_doc_no(name)
        rec["状态"] = status
        recs.append(rec)
    while recs and not any(recs[-1].values()):
        recs.pop()
    return recs


def build_index(xlsx_path: str) -> dict:
    """解析下载的 xlsx，写 preference_index.json，返回统计。"""
    import openpyxl
    # 源文件指纹：整套同步按"内容版本 diff"运作，可产物本身若不记下它是从哪份
    # 字节建的，脱离 state.json 就说不清自己对应哪个版本。这里就地算原始 xlsx 的
    # SHA1 落进索引——纯本地读，不联网。
    with open(xlsx_path, "rb") as fh:
        src_sha1 = hashlib.sha1(fh.read()).hexdigest()
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    data = {}
    for ws in wb.worksheets:
        if "有效" in ws.title:
            data["有效"] = _parse_sheet(ws, "有效")
        elif "失效" in ws.title:
            data["失效"] = _parse_sheet(ws, "失效")
    version = ""
    m = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日", os.path.basename(xlsx_path))
    if not m:
        # 同步器把产物统一存成 latest.xlsx，文件名不带日期；版本日期写在某张表的
        # 表头前几行里，逐格找第一个"YYYY年M月D日"。都找不到才留空。
        for ws in wb.worksheets:
            found = False
            for row in ws.iter_rows(values_only=True, max_row=5):
                for c in row:
                    m2 = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日", _norm(c))
                    if m2:
                        version = f"{m2.group(1)}-{int(m2.group(2)):02d}-{int(m2.group(3)):02d}"
                        found = True
                        break
                if found:
                    break
            if found:
                break
    else:
        version = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    # 文号抽取覆盖率自检：目录里每条政策本都应带文号，抽不到多半是官方那列写法
    # 改版（换括号、加机关前缀、拆两行）。把它当早期回归信号落进索引，sync 时报出，
    # 免得抽取静默退化只能靠将来检索召回下滑才察觉。
    total_recs = sum(len(g) for g in data.values())
    miss_doc = sum(1 for g in data.values() for r in g if not (r.get("文号") or "").strip())
    index = {
        "版本日期": version,
        "源文件SHA1": src_sha1,
        "构建时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "有效条数": len(data.get("有效", [])),
        "失效条数": len(data.get("失效", [])),
        "文号抽取失败": miss_doc,
        "总条数": total_recs,
        "记录": data,
    }
    INDEX_PATH.write_text(json.dumps(index, ensure_ascii=False, indent=1),
                          encoding="utf-8")
    return index


# ── 定位目录直链（喂给 tax_sync 的 locate 回调）────────────────────────────
def _find_xlsx(html: str, base: str):
    """在一页 HTML 里定位《减免税政策代码目录》xlsx 直链，取不到给 None。

    先认「链接文字含目录名 + 落尾是 xlsx」，再退一步认「href 落在目录栏目节点
    （c102373）或路径里带'减免税'」。base 用于把相对 href 拼成绝对直链。
    """
    import urllib.parse
    for m in re.finditer(r"""<a[^>]+href=["']([^"']+)["'][^>]*>(.*?)</a>""",
                         html, re.S | re.I):
        href, inner = m.group(1), re.sub(r"<[^>]+>", "", m.group(2))
        if ANCHOR in inner.replace(" ", "").replace("\n", "") and \
                href.lower().endswith((".xlsx", ".xls")):
            return urllib.parse.urljoin(base, href)
    for m in re.finditer(r"""href=["']([^"']+\.(?:xlsx|xls))["']""", html, re.I):
        href = m.group(1)
        if "c102373" in href or "减免税" in urllib.parse.unquote(href):
            return urllib.parse.urljoin(base, href)
    return None


def locate():
    """按「栏目页 → 首页」顺序定位《减免税政策代码目录》xlsx 直链。

    栏目页是目录的常驻落点；抓取异常或那页改版时回退官网首页再找一次，
    避免单点入口失效就整条同步断掉。全部入口都定位不到才抛错，让同步器
    保留已有索引、不覆盖。
    """
    import tax_http
    ua = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"}
    errors = []
    for page in ENTRY_PAGES:
        try:
            html = tax_http.get(page, headers=ua, timeout=60, verify=True).content.decode(
                "utf-8", "ignore")
        except Exception as e:
            errors.append(f"{page}: {type(e).__name__} {e}")
            continue
        url = _find_xlsx(html, page)
        if url:
            return {"url": url, "ext": ".xlsx"}
        errors.append(f"{page}: 页内未定位到目录 xlsx 直链")
    raise RuntimeError("未能在栏目页/首页定位目录 xlsx 直链：" + " | ".join(errors))


# ── 离线查询 ───────────────────────────────────────────────────────────────
_FIELDS_WEIGHT = {
    "代码": 100, "文号": 40, "减免项目名称": 30,
    "文件标题": 20, "政策名称": 20, "小类": 12,
    "大类": 10, "收入种类": 8, "优惠条款": 3, "关联政策": 3,
}


def _load_index():
    if not INDEX_PATH.exists():
        raise SystemExit(f"[缺索引] 未找到 {INDEX_PATH.name}，先运行：python preference.py sync")
    return json.loads(INDEX_PATH.read_text(encoding="utf-8"))


def _score(rec, words):
    total = 0
    for w in words:
        hit = 0
        for field, weight in _FIELDS_WEIGHT.items():
            if w in (rec.get(field) or ""):
                hit = max(hit, weight)
        if not hit:
            return 0                # AND：任一词不命中即排除
        total += hit
    if rec.get("状态") == "有效":
        total += 5
    return total


def _cited_url(doc_no, cited_path=None) -> str:
    """按目录里的文号取已核实的官方原文链接；取不到给空串。

    文号先经 tax_terms.doc_number_of 归一，与 tax_gov_list.seed_cited_cache
    写入缓存时用的是同一把键——不归一就查不到（缓存里存的是"2026年第19号"，
    目录里写的是"国家税务总局公告2026年第19号"）。纯本地读缓存，不联网。
    """
    if not doc_no:
        return ""
    import tax_terms as TT
    import tax_cited as CITED
    return CITED.get_cited_link(TT.doc_number_of(doc_no), cited_path)


def query(words, code=None, type_=None, status=None, cat=None, limit=10,
          as_json=False, cited_path=None):
    idx = _load_index()
    pool = []
    for st, rows in idx["记录"].items():
        if status and status != st:
            continue
        pool.extend(rows)
    if code:
        pool = [r for r in pool if r.get("代码") == code]
    if type_:
        pool = [r for r in pool if type_ in (r.get("收入种类") or "")]
    if cat:
        pool = [r for r in pool if cat in (r.get("大类") or "")]
    scored = [(r, _score(r, words)) for r in pool] if words else [(r, 0) for r in pool]
    matched = [(r, s) for r, s in scored if s > 0] if words else [
        (r, s) for r, s in scored if code or type_ or cat or status]
    matched.sort(key=lambda x: (-x[1], x[0].get("代码", "")))
    matched = matched[:limit]
    if as_json:
        print(json.dumps({"版本": idx.get("版本日期"), "命中": len(matched),
                          "结果": [r for r, _ in matched]}, ensure_ascii=False, indent=1))
        return matched
    print(f"命中 {len(matched)} 条（版本 {idx.get('版本日期')}）\n")
    for r, _ in matched:
        print(f"[{r.get('状态')}] {r.get('代码')} | {r.get('文件标题')}")
        if r.get("文号"):
            print(f"    文号：{r['文号']}")
            url = _cited_url(r["文号"], cited_path)
            print(f"    官方链接：{url}" if url
                  else "    官方链接：<未缓存，检索原文后可用 tax_gov_list.py seed-cache 预热>")
        print(f"    税种：{r.get('收入种类')} > {r.get('大类')} > {r.get('小类')}")
        if r.get("减免项目名称"):
            print(f"    减免项目：{r['减免项目名称']}")
        if r.get("优惠条款"):
            print(f"    优惠条款：{r['优惠条款']}")
        if r.get("状态") == "失效":
            print(f"    有效期：{r.get('有效期起')} ~ {r.get('有效期止')}")
    return matched


def list_types():
    idx = _load_index()
    from collections import Counter
    c = Counter()
    for rows in idx["记录"].values():
        for r in rows:
            if r.get("收入种类"):
                c[r["收入种类"]] += 1
    print(f"版本 {idx.get('版本日期')} | 有效 {idx.get('有效条数')} / 失效 {idx.get('失效条数')}"
          f" | 文号抽取失败 {idx.get('文号抽取失败', '?')}/{idx.get('总条数', '?')}"
          f" | 源文件SHA1 {str(idx.get('源文件SHA1', '?'))[:8]}")
    for name, n in c.most_common():
        print(f"  {n:>4}  {name}")


# ── CLI ────────────────────────────────────────────────────────────────────
def _sync(check=False, force=False, as_json=False):
    import tax_sync
    syn = tax_sync.Synchronizer("tax-preference", locate, build_index, ext=".xlsx")
    res = syn.sync(check_only=check, force=force)
    # 真正重建过时，把本次 build_index 落进索引的文号抽取覆盖率带回回显（同步时点暴露改版）
    if res.get("动作") == "已更新" and INDEX_PATH.exists():
        try:
            meta = json.loads(INDEX_PATH.read_text(encoding="utf-8"))
            res["文号抽取失败"] = meta.get("文号抽取失败")
            res["总条数"] = meta.get("总条数")
            res["源文件SHA1"] = meta.get("源文件SHA1")
        except Exception:
            pass
    if as_json:
        print(json.dumps(res, ensure_ascii=False, indent=1))
    else:
        print(("[成功] " if res.get("成功") else "[失败] ") + res.get("动作", ""))
        for k in ("本地版本", "远端链接", "新版本", "源文件SHA1", "字节数", "错误"):
            if res.get(k):
                print(f"    {k}：{res[k]}")
        if res.get("动作") == "已更新" and res.get("文号抽取失败") is not None:
            n, tot = res["文号抽取失败"], res.get("总条数")
            flag = "  ← 若该数较上次明显跳增，多半是目录文号列改版，须核对解析" if n else ""
            print(f"    文号抽取失败：{n}/{tot} 条{flag}")
    return 0 if res.get("成功") else 1


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description="《减免税政策代码目录》查询式子技能")
    sub = ap.add_subparsers(dest="cmd")
    p_sync = sub.add_parser("sync", help="检查并按需下载重建索引")
    p_sync.add_argument("--check", action="store_true")
    p_sync.add_argument("--force", action="store_true")
    p_sync.add_argument("--json", action="store_true")
    p_q = sub.add_parser("query", help="离线查询")
    p_q.add_argument("words", nargs="*")
    p_q.add_argument("--code")
    p_q.add_argument("--type", dest="type_")
    p_q.add_argument("--status")
    p_q.add_argument("--cat")
    p_q.add_argument("-n", "--limit", type=int, default=10)
    p_q.add_argument("--json", action="store_true")
    sub.add_parser("list-types", help="列收入种类及条数")
    a = ap.parse_args(argv)

    if a.cmd == "sync":
        return _sync(a.check, a.force, a.json)
    if a.cmd == "query":
        query(a.words, a.code, a.type_, a.status, a.cat, a.limit, a.json)
        return 0
    if a.cmd == "list-types":
        list_types()
        return 0
    ap.print_help()
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
